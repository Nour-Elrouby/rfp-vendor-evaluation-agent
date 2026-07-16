import html
import re
from pathlib import Path

import openpyxl
import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


def validate_file_signature(contents: bytes, extension: str) -> None:
    """Reject obvious extension/content mismatches before invoking parsers."""
    if extension == ".pdf" and not contents.startswith(b"%PDF-"):
        raise ValueError("The uploaded file is not a valid PDF.")
    if extension in {".docx", ".xlsx"} and not contents.startswith(b"PK"):
        raise ValueError(f"The uploaded file is not a valid {extension[1:].upper()} file.")

def read_pdf(file_path: str | Path) -> str:
    pages: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def read_docx(file_path: str | Path) -> str:
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def read_xlsx(file_path: str | Path) -> str:
    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                lines.append(" ".join(str(cell) for cell in row if cell is not None))
        return "\n".join(lines)
    finally:
        workbook.close()


def extract_vendor_text(file_path: str | Path) -> str:
    """Extracts text from a supported proposal and rejects empty documents."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Proposal file does not exist: {path}")

    extension = path.suffix.lower()
    readers = {".pdf": read_pdf, ".docx": read_docx, ".xlsx": read_xlsx}
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension or '<none>'}'. Use: {supported}.")

    text = html.unescape(readers[extension](path))
    text = re.sub(r"\(cid:\d+\)", " ", text)
    if not text.strip():
        raise ValueError("The uploaded proposal contains no extractable text.")
    return text.strip()
